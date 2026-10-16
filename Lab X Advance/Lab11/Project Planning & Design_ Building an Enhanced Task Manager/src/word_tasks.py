import docx
import os
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from utils import setup_logging

logger = setup_logging(__name__)


class WordTasks:
    """Handles advanced Word document operations for the Document Agent."""

    def __init__(self):
        logger.info("WordTasks initialized.")

    def _replace_placeholder_in_paragraph(self, paragraph, placeholder, value):
        """Replaces a single placeholder in a paragraph's runs."""
        if placeholder not in paragraph.text:
            return False  # Placeholder not found in this paragraph

        for run in paragraph.runs:
            if placeholder in run.text:
                run.text = run.text.replace(placeholder, str(value))
                logger.debug(f"Replaced '{placeholder}' with '{value}' in paragraph.")

        # Placeholder split across several runs: rewrite the paragraph in its first run
        if placeholder in paragraph.text and paragraph.runs:
            full_text = paragraph.text.replace(placeholder, str(value))
            for i, run in enumerate(paragraph.runs):
                run.text = full_text if i == 0 else ""
        return True

    @staticmethod
    def _set_table_borders(table):
        """Draws single-line borders on every edge (used when the template has no 'Table Grid' style)."""
        tbl_pr = table._tbl.tblPr
        borders = OxmlElement('w:tblBorders')
        for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
            element = OxmlElement(f'w:{edge}')
            element.set(qn('w:val'), 'single')
            element.set(qn('w:sz'), '4')
            element.set(qn('w:space'), '0')
            element.set(qn('w:color'), '000000')
            borders.append(element)
        tbl_pr.append(borders)

    def _build_table(self, document, table_data):
        cols = len(table_data[0])
        table = document.add_table(rows=0, cols=cols)
        try:
            table.style = 'Table Grid'
        except Exception:
            logger.warning("Style 'Table Grid' not available, drawing borders manually.")
            self._set_table_borders(table)
        for r_idx, row_data in enumerate(table_data):
            row = table.add_row()
            for c_idx, cell_data in enumerate(row_data):
                cell = row.cells[c_idx]
                cell.text = str(cell_data)
                if r_idx == 0 and cell.paragraphs[0].runs:
                    cell.paragraphs[0].runs[0].bold = True
        return table

    def generate_report_from_template(self, template_path, output_path, data_placeholders, table_data=None):
        """
        Generates a Word document from a template, filling placeholders and optional table data.
        `data_placeholders` is a dictionary where keys are placeholders (e.g., '{{REPORT_TITLE}}')
        and values are the replacement strings.
        `table_data` is a list of lists, where the first sub-list is the header row.
        """
        if not os.path.exists(template_path):
            logger.error(f"Word template not found: {template_path}")
            return False

        # Keys may be given bare ("REPORT_TITLE") or wrapped ("{{REPORT_TITLE}}")
        data_placeholders = {(k if k.startswith("{{") else "{{" + k + "}}"): v
                             for k, v in data_placeholders.items()}

        try:
            document = docx.Document(template_path)

            # Replace placeholders in paragraphs
            for paragraph in document.paragraphs:
                for placeholder, value in data_placeholders.items():
                    self._replace_placeholder_in_paragraph(paragraph, placeholder, value)

            # Replace placeholders in table cells
            for table in document.tables:
                for row in table.rows:
                    for cell in row.cells:
                        for paragraph in cell.paragraphs:
                            for placeholder, value in data_placeholders.items():
                                self._replace_placeholder_in_paragraph(paragraph, placeholder, value)

            # Add dynamic table if provided
            if table_data:
                logger.info("Adding dynamic table to the document.")
                anchor = next((p for p in document.paragraphs
                               if "{{TABLE_PLACEHOLDER}}" in p.text), None)
                table = self._build_table(document, table_data)
                if anchor is not None:
                    anchor._element.addprevious(table._tbl)       # put the table where the placeholder is
                    anchor._element.getparent().remove(anchor._element)
                else:
                    logger.warning("Table placeholder '{{TABLE_PLACEHOLDER}}' not found in template. "
                                   "Appending table to end.")

            document.save(output_path)
            logger.info(f"Report generated from template and saved to '{output_path}'.")
            return True

        except Exception as e:
            logger.error(f"Error generating report from template: {e}")
            return False

    def extract_text_from_word(self, docx_path):
        """Extracts all text from a Word document."""
        if not os.path.exists(docx_path):
            logger.error(f"Word document not found: {docx_path}")
            return None

        try:
            document = docx.Document(docx_path)
            full_text = [para.text for para in document.paragraphs]

            # Also extract text from tables
            for table in document.tables:
                for row in table.rows:
                    full_text.append("\t".join(cell.text for cell in row.cells))

            logger.info(f"Text extracted from Word document: {docx_path}")
            return "\n".join(full_text)

        except Exception as e:
            logger.error(f"Error extracting text from Word document '{docx_path}': {e}")
            return None
